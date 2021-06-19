from openpyxl import  load_workbook
from openpyxl.chart import Reference, Series, LineChart
import pandas as pd
import jdatetime
from openpyxl.chart.axis import DateAxis


df = pd.read_excel('wells.xlsx')

df['Date'] = "13" + df.yy.astype(int).astype(str) + "-" + df.mm.astype(int).astype(str) + "-" + df.dd.astype(
    int).astype(str)

df['Date'] = df['Date'].apply(
    lambda x: jdatetime.date(year=int(x.split("-")[0]), month=int(x.split("-")[1]),
                             day=int(x.split("-")[2])).togregorian())
df = df[['Date', 'water_cut']]
df.set_index('Date', inplace=True)
df.to_excel("res.xlsx")

wb = load_workbook('res.xlsx')
ws = wb.active

chart = LineChart()
chart.title = "title"
chart.style = 1
chart.y_axis.title = "water_cut"
chart.y_axis.crossAx = 500
chart.x_axis = DateAxis(crossAx=100)
chart.x_axis.majorTimeUnit = "days"
chart.x_axis.title = "Date"
data = Reference(ws, min_col=2, min_row=1, max_col=2, max_row=40)
chart.add_data(data, titles_from_data=True)
dates = Reference(ws, min_col=1, min_row=2, max_row=30)
chart.set_categories(dates)

ws.add_chart(chart, "D1")

wb.save("res.xlsx")
